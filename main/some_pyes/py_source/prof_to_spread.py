import openpyxl

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sample Sheet"

months = [
    "Jan",
    "Feb",
    "Mar",
    "Apr",
    "May",
    "Jun",
    "Jul",
    "Aug",
    "Sep",
    "Oct",
    "Nov",
    "Dec",
]

for index, month in enumerate(months):
    ws['A' + str(index + 1)] = month
    print(f"Enter profit for {month}:", end=' ')
    profit = input()
    ws['B' + str(index + 1)] = profit

wb.save("py_data/second.xlsx")
